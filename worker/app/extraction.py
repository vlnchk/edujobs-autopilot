"""Deterministic extraction for the source formats supported by the MVP."""

from __future__ import annotations

import re
import unicodedata
import zipfile
from dataclasses import dataclass
from io import BytesIO
from urllib.parse import urljoin, urlsplit

from bs4 import BeautifulSoup
from docx import Document
from docx.document import Document as DocumentType
from docx.opc.constants import RELATIONSHIP_TYPE
from pypdf import PdfReader


class ExtractionError(RuntimeError):
    """Raised when supported content cannot be parsed into usable text."""


class ExtractedContentTooLarge(ExtractionError):
    """Raised rather than truncating extracted content."""


@dataclass(frozen=True)
class ExtractedContent:
    title: str | None
    text: str
    links: list[str]


def _normalize_line(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def _join_blocks(blocks: list[str], max_chars: int) -> str:
    normalized: list[str] = []
    current_length = 0
    for block in blocks:
        clean_block = _normalize_line(block)
        if not clean_block or (normalized and normalized[-1] == clean_block):
            continue
        added_length = len(clean_block) + (2 if normalized else 0)
        if current_length + added_length > max_chars:
            raise ExtractedContentTooLarge(
                "Extracted content exceeds MAX_EXTRACTED_CHARS"
            )
        normalized.append(clean_block)
        current_length += added_length
    return "\n\n".join(normalized)


def _useful_link(base_url: str, href: str) -> str | None:
    href = href.strip()
    if not href or href.startswith("#"):
        return None
    absolute = urljoin(base_url, href)
    if urlsplit(absolute).scheme.lower() not in {"http", "https", "mailto", "tel"}:
        return None
    return absolute


def _unique(values: list[str]) -> list[str]:
    return list(dict.fromkeys(values))


def extract_html(content: bytes, base_url: str, max_chars: int) -> ExtractedContent:
    soup = BeautifulSoup(content, "html.parser")
    title = _normalize_line(soup.title.get_text(" ", strip=True)) if soup.title else None

    for element in soup.find_all(["script", "style", "noscript", "template", "form"]):
        element.decompose()

    root = soup.find("main") or soup.find("article") or soup.body or soup
    chrome_tags = (
        ["nav", "aside"]
        if root.name in {"main", "article"}
        else ["nav", "header", "footer", "aside"]
    )
    for element in root.find_all(chrome_tags):
        element.decompose()
    block_names = {"h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "pre", "blockquote"}
    blocks: list[str] = []
    for element in root.find_all(list(block_names)):
        if any(parent.name in block_names for parent in element.parents if parent is not root):
            continue
        text = element.get_text(" ", strip=True)
        if element.name and element.name.startswith("h"):
            level = int(element.name[1])
            text = f"{'#' * level} {text}"
        elif element.name == "li":
            text = f"- {text}"
        blocks.append(text)

    if not blocks:
        blocks = [line for line in root.get_text("\n").splitlines() if line.strip()]

    links = [
        useful
        for anchor in root.find_all("a", href=True)
        if (useful := _useful_link(base_url, str(anchor["href"]))) is not None
    ]
    return ExtractedContent(
        title=title or None,
        text=_join_blocks(blocks, max_chars),
        links=_unique(links),
    )


def _pdf_links(reader: PdfReader) -> list[str]:
    links: list[str] = []
    for page in reader.pages:
        for annotation_ref in page.get("/Annots", []):
            try:
                annotation = annotation_ref.get_object()
                action = annotation.get("/A")
                uri = action.get("/URI") if action else None
            except (AttributeError, KeyError, TypeError):
                continue
            if isinstance(uri, str) and urlsplit(uri).scheme.lower() in {
                "http",
                "https",
                "mailto",
                "tel",
            }:
                links.append(uri)
    return _unique(links)


def extract_pdf(content: bytes, max_chars: int) -> ExtractedContent:
    try:
        reader = PdfReader(BytesIO(content))
        blocks: list[str] = []
        current_length = 0
        for page in reader.pages:
            block = _normalize_line(page.extract_text() or "")
            if not block or (blocks and blocks[-1] == block):
                continue
            added_length = len(block) + (2 if blocks else 0)
            if current_length + added_length > max_chars:
                raise ExtractedContentTooLarge(
                    "Extracted content exceeds MAX_EXTRACTED_CHARS"
                )
            blocks.append(block)
            current_length += added_length
        text = "\n\n".join(blocks)
        metadata_title = reader.metadata.title if reader.metadata else None
        title = _normalize_line(metadata_title) if metadata_title else None
        links = _pdf_links(reader)
    except ExtractedContentTooLarge:
        raise
    except Exception as exc:
        raise ExtractionError("Unable to parse PDF content") from exc
    return ExtractedContent(title=title, text=text, links=links)


def _docx_uncompressed_size(content: bytes) -> int:
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            return sum(item.file_size for item in archive.infolist())
    except (OSError, zipfile.BadZipFile) as exc:
        raise ExtractionError("Unable to parse DOCX container") from exc


def _docx_links(document: DocumentType) -> list[str]:
    links: list[str] = []
    for relationship in document.part.rels.values():
        if relationship.reltype == RELATIONSHIP_TYPE.HYPERLINK:
            target = relationship.target_ref
            if urlsplit(target).scheme.lower() in {"http", "https", "mailto", "tel"}:
                links.append(target)
    return _unique(links)


def extract_docx(
    content: bytes, max_chars: int, max_uncompressed_bytes: int
) -> ExtractedContent:
    if _docx_uncompressed_size(content) > max_uncompressed_bytes:
        raise ExtractedContentTooLarge("DOCX expands beyond MAX_DOWNLOAD_BYTES")

    try:
        document = Document(BytesIO(content))
        blocks: list[str] = []
        for paragraph in document.paragraphs:
            text = paragraph.text
            if paragraph.style and paragraph.style.name.startswith("Heading"):
                match = re.search(r"(\d+)$", paragraph.style.name)
                level = int(match.group(1)) if match else 1
                text = f"{'#' * min(level, 6)} {text}"
            blocks.append(text)
        for table in document.tables:
            for row in table.rows:
                blocks.append(" | ".join(cell.text for cell in row.cells))

        title_value = document.core_properties.title
        title = _normalize_line(title_value) if title_value else None
        text = _join_blocks(blocks, max_chars)
        links = _docx_links(document)
    except ExtractedContentTooLarge:
        raise
    except Exception as exc:
        raise ExtractionError("Unable to parse DOCX content") from exc
    return ExtractedContent(title=title, text=text, links=links)


def is_docx(content: bytes) -> bool:
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            names = set(archive.namelist())
    except (OSError, zipfile.BadZipFile):
        return False
    return "[Content_Types].xml" in names and "word/document.xml" in names


_NON_CONTENT_PLACEHOLDERS = {
    "loading",
    "loading content",
    "loading document",
    "document is loading",
    "please wait",
    "please wait while the document loads",
    "enable javascript",
    "javascript is required",
    "please enable javascript",
    "загрузка",
    "идет загрузка",
    "идёт загрузка",
    "загрузка документа",
    "выполняется загрузка",
    "подождите",
    "подождите пожалуйста",
    "пожалуйста подождите",
    "מתבצעת טעינה",
    "טעינה",
    "טוען",
    "נא להמתין",
}


def _placeholder_key(text: str) -> str:
    normalized = unicodedata.normalize("NFKC", _normalize_line(text)).casefold()
    normalized = normalized.lstrip("#*-–—• ")
    normalized = re.sub(r"[^\w\s]+", " ", normalized, flags=re.UNICODE)
    return _normalize_line(normalized)


def has_usable_text(text: str) -> bool:
    normalized = _placeholder_key(text)
    return bool(normalized) and normalized not in _NON_CONTENT_PLACEHOLDERS
